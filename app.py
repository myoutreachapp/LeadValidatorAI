"""
Lead Validator — web app version.
Run with: streamlit run app.py
Then open the local web address it prints (usually http://localhost:8501)
"""

import base64
import io
import json
import os
import tempfile
from pathlib import Path

import docx
import pandas as pd
import requests
import streamlit as st
import yaml

from validate_leads import validate, load_tal

st.set_page_config(page_title="MyOutreach Lead Validator", layout="wide")
st.title("MyOutreach Lead Validator")

BASE_TITLES_PATH = Path(__file__).parent / "job_titles_reference.yaml"
with open(BASE_TITLES_PATH) as f:
    BASE_JOB_TITLES = yaml.safe_load(f)


def parse_custom_functions(text):
    """Format, one per line: FunctionName: title one, title two, title three
    A slash within a segment (e.g. "CIO / CISO") means "these are separate
    acceptable titles", not one combined title — so it's split out into two
    reference entries, same as if they'd been comma-separated. Otherwise "CIO"
    and "CISO" never exist as their own matchable titles, only bundled together
    as one odd string that barely resembles any real lead's actual job title."""
    custom = {}
    for line in text.splitlines():
        line = line.strip()
        if not line or ":" not in line:
            continue
        name, titles = line.split(":", 1)
        name = name.strip()
        title_list = []
        for segment in titles.split(","):
            for variant in segment.split("/"):
                v = variant.strip()
                if v:
                    title_list.append(v)
        if name and title_list:
            custom[name] = title_list
    return custom


ICP_EXTRACTION_SYSTEM_PROMPT = """You are extracting structured lead-qualification criteria from a raw B2B marketing contract or ICP document that a user pastes in.

Return ONLY valid JSON (no markdown fences, no preamble, no commentary) in exactly this shape:

{
  "job_functions": {"FunctionName": ["Title 1", "Title 2"]},
  "allowed_countries": ["Country1", "Country2"],
  "permitted_industries": ["Industry1", "Industry2"],
  "cap_per_company": <int or null>,
  "lead_target": <int or null>,
  "deadline": "<yyyy-mm-dd or null>",
  "exclusions_company": ["term1"],
  "exclusions_title": ["term1"],
  "exclusions_industry": ["term1"],
  "watch_time_minutes_min": <int or null>
}

Rules:
- Group job titles by ACTUAL job function / line-of-business (e.g. "Legal", "Operations", "IT", "Finance", "Knowledge Management", "Executive Leadership") — this describes WHAT the person does, not what tier of company they work at.
- The document may organize titles under labels that look like function groupings but are actually company-size or account-tier segments (e.g. "Small-Segment titles", "Mid-Segment titles", "Enterprise titles", "Corporate & Strategic titles"). These are NOT job functions — don't use them as your "job_functions" keys. Company size belongs elsewhere in the ICP (permitted company size), not in job function grouping.
- When titles are organized by tier/segment like that, first pool every title across all tiers, remove duplicates, then classify each one by what the role actually does (e.g. "Managing Partner"/"Solicitor"/"Attorney" -> Legal; "IT Manager"/"Chief Information Officer" -> IT; "Finance Director"/"CFO" -> Finance; "Operations Manager"/"COO" -> Operations; "Head of Knowledge"/"Chief Knowledge Officer" -> Knowledge Management).
- Only use the document's own labels as function names directly when those labels already ARE functional/departmental (e.g. "Data Leaders", "Marketing Leaders", "Legal", "Finance" describe a type of work, not a company-size tier).
- If the document describes geo-based rules (e.g. "EMEA also gets Marketing titles"), just extract the full combined set of functions/titles here — a human will split by region separately afterward.
- Extract every country mentioned as permitted geography, spelled out in full (e.g. "United States" not "US" or "USA"). Permitted geography is often written as a list of REGIONS, where most regions are groups of US states/provinces (e.g. "Southwest: NV, UT, CO, AZ, NM") — those all imply "United States". But watch for a region entry that is itself just a bare country name with no state/province list underneath it (e.g. a line that just says "Canada") — that means the WHOLE country is permitted and must be added to allowed_countries too, even though it doesn't follow the same "region: state-list" pattern as its neighbors. Don't let the surrounding US-state-heavy formatting cause you to overlook a standalone country line.
- For "permitted_industries": this field is often a comma-separated list where SOME entries have parenthetical examples of sub-categories, e.g. "Finance Orgs (Banks, Investment Banks, Credit Card companies, Fintech), Retail, ... SLED(Fed Gov, Local/State, Higher Ed)". Treat each parenthetical group as ONE top-level industry with its examples, not separate industries — split only on the top-level commas OUTSIDE parentheses. So that example yields industries like "Finance Orgs (Banks, Investment Banks, Credit Card companies, Fintech)", "Retail", ..., "SLED(Fed Gov, Local/State, Higher Ed)" — five-ish entries, not a dozen fragments.
- Dates are often written imprecisely (e.g. "End of September 2026", "Q3 2026", "July-September 2026"). Resolve these to a specific yyyy-mm-dd using reasonable interpretation — "end of September 2026" -> "2026-09-30", "Q3 2026" -> "2026-09-30" (use the END of a stated period for "deadline", since that's the field's purpose). Only use null if literally no date or timeframe is mentioned anywhere.
- The document is very often a labeled table (row label -> value), such as: Delivery Template, Permitted Industries, Permitted Geographic Regions, Company Size, Job Function of Prospect, Level of Responsibility, Term of the Campaign, Lead Delivery Dates/Pacing, Cap per Company, Other Requirements. Go through EVERY labeled row methodically and populate the matching JSON field — don't skip a row just because its value needs light interpretation (a date range, a parenthetical list, a number embedded in a sentence).
- If a field isn't present in the text, use null (for scalars) or an empty list/object (for lists/objects) — never invent data that isn't in the text.
- "cap_per_company" is the max number of leads allowed from one company/account, if stated — look for a row literally labeled "Cap per Company" or similar and extract its exact number.
- Keep "exclusions_company" (company name), "exclusions_title" (job title), and "exclusions_industry" (industry/sector) separate — an instruction to exclude "VC firms" or "PE firms" is an INDUSTRY exclusion, not a company name exclusion, even though the document may not label it explicitly as "industry."
- "permitted_industries" is the list of industries/sectors the ICP targets (the inclusion counterpart to exclusions_industry) — leave it empty if the document doesn't name specific industries (e.g. if industry is entirely implied by an attached account list instead).
"""


def extract_docx_text(file_bytes):
    doc = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def build_content_blocks(icp_text, uploaded_file):
    """Builds the Anthropic API 'content' list from pasted text and/or an
    uploaded screenshot/PDF/docx. Images and PDFs go straight to Claude as-is
    (vision handles reading tables/screenshots); docx text is extracted first
    since the API doesn't accept that format directly."""
    blocks = []
    combined_text = icp_text or ""

    if uploaded_file is not None:
        name = uploaded_file.name.lower()
        data = uploaded_file.getvalue()
        if name.endswith((".png", ".jpg", ".jpeg")):
            media_type = "image/png" if name.endswith(".png") else "image/jpeg"
            blocks.append({
                "type": "image",
                "source": {"type": "base64", "media_type": media_type, "data": base64.b64encode(data).decode()},
            })
        elif name.endswith(".pdf"):
            blocks.append({
                "type": "document",
                "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode()},
            })
        elif name.endswith(".docx"):
            combined_text = (combined_text + "\n" + extract_docx_text(data)).strip()

    if combined_text.strip():
        blocks.append({"type": "text", "text": combined_text})
    return blocks


def extract_icp_with_ai(content_blocks, api_key, workspace_id=None):
    headers = {
        "x-api-key": api_key,
        "anthropic-version": "2023-06-01",
        "content-type": "application/json",
    }
    if workspace_id:
        headers["anthropic-workspace-id"] = workspace_id
    r = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers=headers,
        json={
            "model": "claude-sonnet-5",
            "max_tokens": 2000,
            "system": ICP_EXTRACTION_SYSTEM_PROMPT,
            "messages": [{"role": "user", "content": content_blocks}],
        },
        timeout=60,
    )
    if not r.ok:
        raise RuntimeError(f"{r.status_code} {r.reason}: {r.text}")
    data = r.json()
    text = "".join(b["text"] for b in data["content"] if b.get("type") == "text").strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text.lower().startswith("json"):
            text = text[4:]
    return json.loads(text)


def classify_values_with_ai(unique_values, criteria, api_key, workspace_id, field_label, mode="exclude"):
    """One call, run once per file: classifies every DISTINCT raw value against plain-English
    criteria, catching naming variations without needing every phrasing hand-listed.
    mode='exclude' -> returns values that MATCH the criteria (to reject).
    mode='include' -> returns values that MATCH the criteria (to permit) — same classification,
    just phrased for the inclusion-list use case so the prompt reads naturally either way."""
    if not unique_values or not criteria:
        return []
    verb = "should be excluded" if mode == "exclude" else "match the permitted list and should be allowed"
    prompt = (
        f"Here is a list of real, distinct {field_label} values from a lead database:\n"
        f"{json.dumps(unique_values)}\n\n"
        f"Here are the criteria in plain English: {json.dumps(criteria)}\n\n"
        f"Return ONLY a JSON array containing the EXACT strings (copied verbatim from the first list) "
        f"that {verb} — matching different names, spellings, or phrasings for the same concept "
        f"(e.g. a criterion like 'venture capital' should also match 'VC', 'Venture Capital Firms', "
        f"'Investment Services - VC', etc. if those appear in the list; a criterion like "
        f"'Software/Technology' should also match 'SaaS' or 'Enterprise Software'). "
        f"Return an empty array [] if nothing matches. No commentary, no markdown, just the JSON array."
    )
    headers = {"x-api-key": api_key, "anthropic-version": "2023-06-01", "content-type": "application/json"}
    if workspace_id:
        headers["anthropic-workspace-id"] = workspace_id
    r = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers=headers,
        json={"model": "claude-sonnet-5", "max_tokens": 1500, "messages": [{"role": "user", "content": prompt}]},
        timeout=60,
    )
    if not r.ok:
        raise RuntimeError(f"{r.status_code} {r.reason}: {r.text}")
    data = r.json()
    text = "".join(b["text"] for b in data["content"] if b.get("type") == "text").strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text.lower().startswith("json"):
            text = text[4:]
    return json.loads(text)


def classify_tal_aliases_with_ai(failed_leads, tal_entries, api_key, workspace_id):
    """Uses actual knowledge of real companies (not string similarity) to catch
    cases like 'McDonald's' / 'mcdonalds.co.uk' failing to match a TAL entry for
    'McDonalds Ltd' / 'mcdonalds.com' — same real company, different legal-entity
    name or country-specific domain. Only matches on genuine confidence, not
    superficial name/domain resemblance (a real risk with plain fuzzy matching:
    e.g. 'Acme Corp' vs 'Acme Holdings Inc' can look similar but be unrelated)."""
    if not failed_leads or not tal_entries:
        return []
    prompt = (
        f"Here is a list of leads that FAILED to match a Target Account List (TAL) by exact domain:\n"
        f"{json.dumps(failed_leads)}\n\n"
        f"Here is the full TAL (the approved account list), by company name and domain:\n"
        f"{json.dumps(tal_entries)}\n\n"
        f"For each failed lead, determine whether it is actually THE SAME REAL-WORLD COMPANY as one "
        f"of the TAL entries — just under a different domain (e.g. a country-specific TLD like "
        f"'mcdonalds.fr' vs 'mcdonalds.com') or a different legal-entity name variant (e.g. 'McDonald's' "
        f"vs 'McDonalds Ltd' vs 'McDonalds Restaurants'). Base this on your actual knowledge of real "
        f"companies, NOT on superficial spelling similarity — two differently-named companies that merely "
        f"look alike (e.g. 'Acme Corp' vs 'Acme Holdings Inc') are NOT a match unless you have genuine "
        f"reason to believe they're the same organization. When genuinely unsure, do not include it.\n\n"
        f"Return ONLY a JSON array of confirmed matches, in this exact shape:\n"
        f'[{{"failed_domain": "...", "matched_tal_company": "...", "matched_tal_domain": "..."}}]\n'
        f"Return an empty array [] if nothing matches. No commentary, no markdown, just the JSON array."
    )
    headers = {"x-api-key": api_key, "anthropic-version": "2023-06-01", "content-type": "application/json"}
    if workspace_id:
        headers["anthropic-workspace-id"] = workspace_id
    r = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers=headers,
        json={"model": "claude-sonnet-5", "max_tokens": 2000, "messages": [{"role": "user", "content": prompt}]},
        timeout=60,
    )
    if not r.ok:
        raise RuntimeError(f"{r.status_code} {r.reason}: {r.text}")
    data = r.json()
    text = "".join(b["text"] for b in data["content"] if b.get("type") == "text").strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text.lower().startswith("json"):
            text = text[4:]
    return json.loads(text)


STANDARD_FIELDS = [
    "first_name", "last_name", "email", "phone", "job_title", "company",
    "domain", "industry", "employee_size", "country", "state",
    "watch_time_minutes",
]

COMMON_COUNTRIES = [
    "United States", "Canada", "Mexico", "United Kingdom", "Ireland", "France",
    "Germany", "Netherlands", "Sweden", "Spain", "Italy", "Australia",
    "New Zealand", "India", "Singapore", "United Arab Emirates", "Brazil",
    "Cyprus", "Kazakhstan",
]

# ---------------------------------------------------------------------------
# Sidebar — API keys (kept only in this session, never saved to disk)
# ---------------------------------------------------------------------------
def get_secret(name, default=""):
    try:
        return st.secrets.get(name, default)
    except Exception:
        return default


with st.sidebar:
    st.header("Settings")

    # Keys set by the admin in secrets.toml are used SILENTLY — never rendered
    # into any input box, masked or not, since a value sent to the browser to
    # pre-fill a field can be read by that user via dev tools regardless of
    # type="password". A visible box only appears as a fallback for someone to
    # type their OWN key when no admin-configured one exists.
    zb_secret = get_secret("ZEROBOUNCE_API_KEY")
    if zb_secret:
        zb_key = zb_secret
        os.environ["ZEROBOUNCE_API_KEY"] = zb_key
        st.caption("ZeroBounce: using the key configured by your admin.")
    else:
        zb_key = st.text_input("ZeroBounce API key", type="password",
                                help="No admin key configured — enter your own for this session.")
        if zb_key:
            os.environ["ZEROBOUNCE_API_KEY"] = zb_key
    st.caption("Leave blank to skip email validity checking for this run.")

    anthropic_secret = get_secret("ANTHROPIC_API_KEY")
    if anthropic_secret:
        anthropic_key = anthropic_secret
        st.caption("Anthropic: using the key configured by your admin.")
    else:
        anthropic_key = st.text_input("Anthropic API key (for AI ICP extraction)", type="password",
                                       help="No admin key configured — enter your own for this session.")

    anthropic_workspace_id = get_secret("ANTHROPIC_WORKSPACE_ID") or st.text_input(
        "Anthropic workspace ID (only needed if you see a 'not scoped to a workspace' error)",
        help="Find it in the Console under the workspace you want to use — starts with wrkspc_.")

# ---------------------------------------------------------------------------
# Step 0 — paste raw ICP/contract text, let AI extract structured criteria
# ---------------------------------------------------------------------------
st.header("0. Paste or upload your ICP from the contract (optional — AI-assisted)")
st.caption("Paste the raw ICP text, or upload a screenshot, PDF, or the contract's .docx file. "
           "Claude will pull out job functions, countries, cap per company, exclusions, etc. and "
           "pre-fill the form below. You can still edit anything it gets wrong.")
icp_text = st.text_area("Raw ICP / contract text (optional if uploading a file below)", height=120,
                         placeholder="Paste the ICP table, job titles list, permitted regions, etc. here...")
icp_file = st.file_uploader("...or upload a screenshot, PDF, or .docx of the contract",
                             type=["png", "jpg", "jpeg", "pdf", "docx"], key="icp_file")

def merge_icp_extractions(old, new):
    """Merges a new extraction into whatever was already captured, instead of
    overwriting wholesale — so uploading the contract in parts (or re-running
    extraction) accumulates fields rather than losing earlier ones."""
    merged = dict(old) if old else {}
    for key, value in (new or {}).items():
        if key == "job_functions":
            existing = dict(merged.get("job_functions") or {})
            for fn, titles in (value or {}).items():
                existing[fn] = list(dict.fromkeys((existing.get(fn) or []) + (titles or [])))
            merged["job_functions"] = existing
        elif isinstance(value, list):
            existing = merged.get(key) or []
            merged[key] = list(dict.fromkeys(existing + (value or [])))
        else:
            # scalar (cap_per_company, lead_target, deadline, etc.) — keep the new
            # value if this extraction actually found one; otherwise don't clobber
            # a value an earlier extraction already found.
            if value not in (None, ""):
                merged[key] = value
    return merged


if st.button("Extract with AI"):
    if not anthropic_key:
        st.error("Add your Anthropic API key in the sidebar first.")
    elif not icp_text.strip() and not icp_file:
        st.error("Paste some ICP text or upload a file first.")
    else:
        try:
            with st.spinner("Reading the ICP..."):
                blocks = build_content_blocks(icp_text, icp_file)
                new_icp = extract_icp_with_ai(blocks, anthropic_key, anthropic_workspace_id)
                st.session_state["ai_icp"] = merge_icp_extractions(st.session_state.get("ai_icp", {}), new_icp)
            st.success("Extracted and merged with anything already captured — form below is pre-filled. Check it over before running.")
        except Exception as e:
            st.error(f"AI extraction failed: {e}")

if st.session_state.get("ai_icp") and st.button("Clear extracted ICP and start over"):
    st.session_state["ai_icp"] = {}
    st.rerun()

ai_icp = st.session_state.get("ai_icp", {})
if ai_icp:
    with st.expander("What the AI extracted (click to review raw output)"):
        st.json(ai_icp)

# fold any AI-found countries into the selectable options so they're valid defaults
ai_countries = ai_icp.get("allowed_countries") or []
for c in ai_countries:
    if c not in COMMON_COUNTRIES:
        COMMON_COUNTRIES.append(c)

# ---------------------------------------------------------------------------
# Step 1 — lead file
# ---------------------------------------------------------------------------
st.header("1. Upload the lead file")
leads_file = st.file_uploader("Lead file (.xlsx, .xls, or .csv)", type=["xlsx", "xls", "csv"])

leads_df = None
if leads_file:
    leads_df = pd.read_csv(leads_file) if leads_file.name.endswith("csv") else pd.read_excel(leads_file)
    st.success(f"Loaded {len(leads_df)} rows. Columns found: {', '.join(leads_df.columns)}")

# ---------------------------------------------------------------------------
# Step 1b — add job functions/titles this campaign needs that aren't in the
# shared reference file yet (e.g. Legal, Finance, HR — anything beyond the
# default Data/Product/Marketing personas). Pre-filled from AI extraction above.
# ---------------------------------------------------------------------------
st.header("1b. Add any new job functions for this campaign (optional)")
st.caption(
    "Only needed if this client's ICP targets a persona not already in the shared list "
    f"({', '.join(BASE_JOB_TITLES.keys())}). One function per line, format: "
    "`FunctionName: title one, title two, title three`. Auto-filled if you used AI extraction above."
)
ai_functions_default = ""
combined_functions_for_default = {**(st.session_state.get("loaded_custom_titles") or {}), **(ai_icp.get("job_functions") or {})}
if combined_functions_for_default:
    ai_functions_default = "\n".join(
        f"{name}: {', '.join(titles)}" for name, titles in combined_functions_for_default.items()
    )
custom_functions_text = st.text_area(
    "New job functions and titles",
    value=ai_functions_default,
    placeholder="Legal: General Counsel, VP Legal, Chief Legal Officer, Head of Compliance\n"
                "Finance: CFO, VP Finance, Controller",
    height=100,
)
custom_functions = parse_custom_functions(custom_functions_text)
ALL_JOB_TITLES = {**BASE_JOB_TITLES, **custom_functions}
if custom_functions:
    st.success(f"Added: {', '.join(custom_functions.keys())} — now selectable below.")

# ---------------------------------------------------------------------------
# Step 2 — campaign setup: load a saved config, or build a new one
# ---------------------------------------------------------------------------
def saved_config_to_ai_icp_shape(config):
    """Converts an uploaded saved-config YAML into the same shape used for
    AI-extracted ICPs, so it can reuse all the same pre-fill wiring in the
    builder form below instead of duplicating it."""
    job_functions = {}
    for bucket in (config.get("geo_function_rules") or {}).values():
        for fn in bucket.get("allowed_functions", []):
            job_functions[fn] = ALL_JOB_TITLES.get(fn, [])
    exclusions = config.get("exclusions") or {}
    return {
        "job_functions": job_functions,
        "allowed_countries": config.get("allowed_countries") or [],
        "permitted_industries": config.get("permitted_industries") or [],
        "cap_per_company": config.get("cap_per_company"),
        "lead_target": (config.get("campaign") or {}).get("lead_target"),
        "deadline": (config.get("campaign") or {}).get("term_end"),
        "exclusions_company": exclusions.get("company_name_contains") or [],
        "exclusions_title": exclusions.get("job_title_contains") or [],
        "exclusions_industry": exclusions.get("industry_contains") or [],
        "watch_time_minutes_min": config.get("watch_time_minutes_min"),
    }


st.header("2. Campaign setup")
mode = st.radio("Do you have a saved config for this campaign already?",
                 ["Upload a saved config (.yaml)", "Build a new one"], horizontal=True,
                 key="setup_mode")

config = None
tal_file = None
use_ai_industry_matching = True
exclude_industries = ""
use_ai_industry_inclusion_matching = True
permitted_industries_text = ""
use_ai_tal_alias_matching = True
tal_required = True

if mode == "Upload a saved config (.yaml)":
    config_upload = st.file_uploader("Campaign config (.yaml)", type=["yaml", "yml"])
    tal_file = st.file_uploader("TAL file for this campaign (.xlsx, .xls, or .csv)",
                                 type=["xlsx", "xls", "csv"])
    if config_upload:
        config = yaml.safe_load(config_upload)
        st.success(f"Loaded config for: {config.get('campaign', {}).get('name', 'this campaign')}")

        # Auto-restore the bundled title list into section 1b, once per upload —
        # this is the fix for configs saved before title-bundling existed, and the
        # normal path for any config saved after it: the title list is NOT part of
        # the campaign settings alone, so without this it silently stays empty and
        # every lead fails with "no titles to match against."
        if config.get("custom_job_titles") and st.session_state.get("_titles_loaded_from") != config_upload.name:
            st.session_state["loaded_custom_titles"] = config["custom_job_titles"]
            st.session_state["_titles_loaded_from"] = config_upload.name
            st.rerun()
        elif not config.get("custom_job_titles"):
            st.warning(
                "This config file doesn't have a bundled title list (it was likely saved before "
                "that feature existed). Make sure section 1b above still has the right titles typed "
                "in for this campaign — otherwise every lead will fail with 'no titles to match against.'"
            )

        with st.expander("Preview this config before running", expanded=True):
            geo_rules = config.get("geo_function_rules") or {}
            all_functions = sorted({fn for b in geo_rules.values() for fn in b.get("allowed_functions", [])})
            exclusions = config.get("exclusions") or {}
            st.markdown(f"""
- **Lead target:** {(config.get('campaign') or {}).get('lead_target', '(not set)')}
- **Deadline:** {(config.get('campaign') or {}).get('term_end', '(not set)')}
- **Cap per company:** {config.get('cap_per_company', '(not set)')}
- **Minimum seniority:** {config.get('minimum_seniority') or '(no minimum)'}
- **Permitted countries:** {', '.join(config.get('allowed_countries') or []) or '(none — global, no restriction)'}
- **Permitted industries:** {', '.join(config.get('permitted_industries') or []) or '(none — no restriction)'}
- **Permitted job functions:** {', '.join(all_functions) or '(none configured)'}
- **Exclude — company name:** {', '.join(exclusions.get('company_name_contains') or []) or '(none)'}
- **Exclude — job title:** {', '.join(exclusions.get('job_title_contains') or []) or '(none)'}
- **Exclude — industry:** {', '.join(exclusions.get('industry_contains') or []) or '(none)'}
- **TAL required:** {config.get('tal_required', True)}
""")
            st.caption("Raw config file, for full detail:")
            st.json(config)

        if st.button("Load into the editable form below to review/adjust before running"):
            st.session_state["ai_icp"] = merge_icp_extractions(
                st.session_state.get("ai_icp", {}), saved_config_to_ai_icp_shape(config))
            st.session_state["loaded_campaign_name"] = (config.get("campaign") or {}).get("name", "New Campaign")
            st.session_state["loaded_minimum_seniority"] = config.get("minimum_seniority")
            st.session_state["loaded_tal_required"] = config.get("tal_required", True)
            st.session_state["setup_mode"] = "Build a new one"
            st.rerun()

else:
    st.caption("Fill this in once per new campaign — or paste the ICP above and let AI pre-fill it. "
               "Download it at the end to reuse next time.")

    st.markdown("**Campaign basics**")
    col1, col2 = st.columns(2)
    with col1:
        campaign_name = st.text_input("Campaign name", st.session_state.get("loaded_campaign_name", "New Campaign"))
        lead_target = st.number_input("Lead target", min_value=0,
                                       value=ai_icp.get("lead_target") or 100)
    with col2:
        term_end = st.text_input("Deadline (yyyy-mm-dd)", ai_icp.get("deadline") or "")
        watch_time_min = st.number_input("Minimum watch time (minutes, 0 = don't check)", min_value=0,
                                          value=ai_icp.get("watch_time_minutes_min") or 0)

    st.markdown("**Geography**")
    default_countries = [c for c in ai_countries if c in COMMON_COUNTRIES] or ["United States"]
    allowed_countries = st.multiselect("Permitted countries (leave blank for global — no restriction)",
                                        COMMON_COUNTRIES, default=default_countries)

    st.markdown("**Industry**")
    permitted_industries_text = st.text_input(
        "Permitted industries (comma-separated, leave blank for no restriction)",
        ", ".join(ai_icp.get("permitted_industries") or []))
    permitted_industries_file = st.file_uploader(
        "...or upload a permitted-industries list (.txt, .csv, or .xlsx — one per line/row)",
        type=["txt", "csv", "xlsx"], key="perm_ind_file")
    use_ai_industry_inclusion_matching = st.checkbox(
        "Use AI to catch different names/spellings for permitted industries",
        value=True,
        help="One AI call, run once on this file, checks every distinct industry value actually "
             "in your data against your permitted list and expands it to the exact matching "
             "spellings/variants found — e.g. a permitted list of 'Software/Technology' also "
             "matching 'SaaS' or 'Enterprise Software' if those appear in the file. "
             "Needs your Anthropic key in the sidebar.")
    exclude_industries = st.text_input("Exclude if industry contains (comma-separated)",
                                        ", ".join(ai_icp.get("exclusions_industry") or []))
    st.caption("Use this for industry-based exclusions (e.g. VC/PE firms) — NOT company name. "
               "Since data sources label these inconsistently, list every variant you've seen: "
               "e.g. 'venture capital, private equity, asset management, investment services'.")
    exclude_industries_file = st.file_uploader("...or upload an industry-exclusion list (.txt, .csv, or .xlsx — one per line/row)",
                                                type=["txt", "csv", "xlsx"], key="excl_ind_file")
    use_ai_industry_matching = st.checkbox(
        "Use AI to catch different names/spellings for these industry exclusions",
        value=True,
        help="Same idea, in reverse: catches variants like 'VC' or 'Investment Services' matching "
             "a 'venture capital' exclusion. Needs your Anthropic key in the sidebar.")

    st.markdown("**Job function & seniority**")
    function_options = list(ALL_JOB_TITLES.keys())
    default_functions = [f for f in ai_icp.get("job_functions", {}).keys() if f in function_options] \
        or (function_options[:2] if len(function_options) >= 2 else function_options)
    allowed_functions = st.multiselect("Permitted job functions", function_options, default=default_functions)
    exclude_titles = st.text_input("Exclude if job title contains (comma-separated)",
                                    ", ".join(ai_icp.get("exclusions_title") or []))
    exclude_titles_file = st.file_uploader("...or upload a title-exclusion list (.txt, .csv, or .xlsx — one per line/row)",
                                            type=["txt", "csv", "xlsx"], key="excl_title_file")
    seniority_options = ["(no minimum)", "Manager", "Director", "VP", "SVP", "EVP", "Chief"]
    loaded_seniority = st.session_state.get("loaded_minimum_seniority")
    seniority_default_idx = seniority_options.index(loaded_seniority) if loaded_seniority in seniority_options else 0
    minimum_seniority_choice = st.selectbox(
        "Minimum seniority (checked separately from job function)", seniority_options,
        index=seniority_default_idx,
        help="E.g. 'IT function, Manager or above' — a lead can match the right function "
             "but still fail here if the title reads as below this level (Engineer, Analyst, "
             "Lead, Specialist, etc. with no Manager/Director/VP+ keyword).")

    with st.expander("Advanced: different job functions allowed by region"):
        st.caption("Only fill this in if, e.g., EMEA allows Marketing titles but North America doesn't.")
        split_geo = st.checkbox("Yes, split by region")
        region_a_countries, region_a_functions = [], []
        region_b_countries, region_b_functions = [], []
        if split_geo:
            region_a_countries = st.multiselect("Region A countries", allowed_countries, key="ra")
            region_a_functions = st.multiselect("Region A allowed functions", function_options, key="raf")
            region_b_countries = st.multiselect("Region B countries", allowed_countries, key="rb")
            region_b_functions = st.multiselect("Region B allowed functions", function_options, key="rbf")

    st.markdown("**Company**")
    exclude_companies = st.text_input("Exclude if company name contains (comma-separated)",
                                       ", ".join(ai_icp.get("exclusions_company") or []))
    exclude_companies_file = st.file_uploader("...or upload an exclusion list (.txt, .csv, or .xlsx — one per line/row)",
                                               type=["txt", "csv", "xlsx"], key="excl_co_file")
    cap_per_company = st.number_input("Max leads per company", min_value=1,
                                       value=ai_icp.get("cap_per_company") or 5)

    st.markdown("**Target Account List (TAL)**")
    tal_file = st.file_uploader("TAL file for this campaign (.xlsx, .xls, or .csv) — one file only",
                                 type=["xlsx", "xls", "csv"])
    tal_sheet = None
    if tal_file:
        if tal_file.name.lower().endswith(".csv"):
            st.caption("CSV detected — no sheet selection needed.")
        else:
            tal_xl = pd.ExcelFile(tal_file)
            tal_sheet = st.selectbox("Which sheet in the TAL has the account list?", tal_xl.sheet_names)

    tal_options = ["Required — reject anything not on it", "Priority only — prefer these accounts, but don't block others"]
    loaded_tal_required = st.session_state.get("loaded_tal_required", True)
    tal_default_idx = 0 if loaded_tal_required else 1
    tal_mode = st.radio(
        "How should the TAL be used?", tal_options, index=tal_default_idx,
        help="Most campaigns require the TAL. Use 'Priority only' when the contract says something like "
             "'prioritize TAL before targeting demographics' — meaning the TAL guides sourcing, but a lead "
             "that isn't on it can still pass if it meets the other ICP criteria (industry, geo, etc.).")
    tal_required = tal_mode.startswith("Required")

    use_ai_tal_alias_matching = st.checkbox(
        "Use AI to catch TAL companies listed under a different domain/name",
        value=True,
        help="Uses actual knowledge of real companies (not spelling similarity) to catch cases like "
             "'JPMC'/'jpmorgan.com' matching a TAL entry for 'JPMorgan Chase & Co'/'jpmorganchase.com', "
             "or a global brand's country-specific domain (mcdonalds.fr vs mcdonalds.com). Runs once per "
             "file, only on leads that failed the exact TAL match. Needs your Anthropic key in the sidebar. "
             "Note: this is knowledge-based, not string-similarity-based — it will NOT loosely accept "
             "unrelated companies that merely look similar.")

    if split_geo:
        geo_function_rules = {
            "Region A": {"countries": region_a_countries, "allowed_functions": region_a_functions},
            "Region B": {"countries": region_b_countries, "allowed_functions": region_b_functions},
        }
    else:
        geo_function_rules = {
            "All": {"countries": allowed_countries, "allowed_functions": allowed_functions}
        }

    def parse_list_file(uploaded):
        if not uploaded:
            return []
        if uploaded.name.lower().endswith(".xlsx"):
            col0 = pd.read_excel(uploaded).iloc[:, 0]
            return [str(v).strip() for v in col0.dropna()]
        text = uploaded.getvalue().decode("utf-8", errors="ignore")
        return [line.strip() for line in text.splitlines() if line.strip()]

    company_exclusions = [t.strip() for t in exclude_companies.split(",") if t.strip()] + parse_list_file(exclude_companies_file)
    title_exclusions = [t.strip() for t in exclude_titles.split(",") if t.strip()] + parse_list_file(exclude_titles_file)
    industry_exclusions = [t.strip() for t in exclude_industries.split(",") if t.strip()] + parse_list_file(exclude_industries_file)
    permitted_industries = [t.strip() for t in permitted_industries_text.split(",") if t.strip()] + parse_list_file(permitted_industries_file)

    config = {
        "campaign": {"name": campaign_name, "lead_target": lead_target, "term_end": term_end},
        "tal": {"file": tal_file.name if tal_file else None, "sheet": tal_sheet, "match_on": "domain"},
        "tal_required": tal_required,
        "allowed_countries": allowed_countries,
        "permitted_industries": permitted_industries,
        "geo_function_rules": geo_function_rules,
        "custom_job_titles": custom_functions,  # bundled in so uploading this config later restores the title list too
        "cap_per_company": cap_per_company,
        "minimum_seniority": None if minimum_seniority_choice == "(no minimum)" else minimum_seniority_choice,
        "exclusions": {
            "company_name_contains": company_exclusions,
            "job_title_contains": title_exclusions,
            "industry_contains": industry_exclusions,
        },
        "qualifying_questions": {},
        "watch_time_minutes_min": watch_time_min,
        "column_mapping": {},  # filled in step 3 below
    }

# ---------------------------------------------------------------------------
# Step 3 — column mapping (auto-guess, let user fix)
# ---------------------------------------------------------------------------
if leads_df is not None and config is not None:
    st.header("3. Match your file's columns")
    st.caption("Auto-matched where possible — check these are right, especially email/company/domain.")
    cols = list(leads_df.columns)
    guessed_mapping = config.get("column_mapping") or {}
    final_mapping = {}
    mcols = st.columns(3)
    for i, field in enumerate(STANDARD_FIELDS):
        with mcols[i % 3]:
            guess = guessed_mapping.get(field)
            if not guess:
                for c in cols:
                    if field.replace("_", "") in c.lower().replace("_", "").replace(" ", ""):
                        guess = c
                        break
            options = ["(none)"] + cols
            default_idx = options.index(guess) if guess in options else 0
            selected = st.selectbox(field, options, index=default_idx, key=f"map_{field}")
            final_mapping[field] = None if selected == "(none)" else selected
    config["column_mapping"] = final_mapping

# ---------------------------------------------------------------------------
# Step 4 — run
# ---------------------------------------------------------------------------
st.header("4. Run validation")
missing_reasons = []
if leads_df is None:
    missing_reasons.append("upload a lead file (Step 1)")
if config is None:
    missing_reasons.append("build or upload a campaign config (Step 2)")
if missing_reasons:
    st.warning("Can't validate yet — you still need to: " + "; ".join(missing_reasons) + ".")

if st.button("Validate leads", type="primary", disabled=(leads_df is None or config is None)):
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        leads_path = tmp / leads_file.name
        leads_path.write_bytes(leads_file.getvalue())

        if use_ai_industry_matching and exclude_industries.strip():
            industry_col = (config.get("column_mapping") or {}).get("industry")
            if not industry_col:
                st.warning("No column is mapped to 'industry' above, so AI industry exclusion matching was skipped.")
            elif not anthropic_key:
                st.warning("Add your Anthropic API key in the sidebar to use AI industry matching. Skipped for this run.")
            else:
                unique_industries = leads_df[industry_col].dropna().unique().tolist()
                criteria = [t.strip() for t in exclude_industries.split(",") if t.strip()]
                try:
                    with st.spinner(f"Checking {len(unique_industries)} distinct industry values against your exclusion criteria..."):
                        matched = classify_values_with_ai(unique_industries, criteria, anthropic_key, anthropic_workspace_id, "industry", mode="exclude")
                    if matched:
                        st.info(f"AI matched these industry values to your EXCLUSION criteria: {', '.join(matched)}")
                        config["exclusions"]["industry_contains"] = list(set(config["exclusions"].get("industry_contains", []) + matched))
                    else:
                        st.caption("AI industry exclusion matching found no additional matches beyond your typed list.")
                except Exception as e:
                    st.error(f"AI industry exclusion matching failed, continuing with your typed exclusion list only: {e}")

        if use_ai_industry_inclusion_matching and permitted_industries_text.strip():
            industry_col = (config.get("column_mapping") or {}).get("industry")
            if not industry_col:
                st.warning("No column is mapped to 'industry' above, so AI permitted-industry matching was skipped.")
            elif not anthropic_key:
                st.warning("Add your Anthropic API key in the sidebar to use AI industry matching. Skipped for this run.")
            else:
                unique_industries = leads_df[industry_col].dropna().unique().tolist()
                criteria = [t.strip() for t in permitted_industries_text.split(",") if t.strip()]
                try:
                    with st.spinner(f"Checking {len(unique_industries)} distinct industry values against your permitted list..."):
                        matched = classify_values_with_ai(unique_industries, criteria, anthropic_key, anthropic_workspace_id, "industry", mode="include")
                    if matched:
                        st.info(f"AI matched these industry values to your PERMITTED criteria: {', '.join(matched)}")
                        config["permitted_industries"] = list(set(config.get("permitted_industries", []) + matched))
                    else:
                        st.caption("AI permitted-industry matching found no additional matches beyond your typed list.")
                except Exception as e:
                    st.error(f"AI permitted-industry matching failed, continuing with your typed list only: {e}")

        config_path = tmp / "config.yaml"
        config_path.write_text(yaml.dump(config))

        # Write the merged titles (shared file + any custom functions added above)
        # so this run's matching includes them, without touching the shared file.
        titles_path = tmp / "job_titles_merged.yaml"
        titles_path.write_text(yaml.dump(ALL_JOB_TITLES))

        tal_dir = tmp
        if tal_file:
            (tmp / tal_file.name).write_bytes(tal_file.getvalue())

        output_path = tmp / "result.xlsx"

        try:
            with st.spinner("Validating..."):
                validate(str(leads_path), str(config_path), str(titles_path), str(tal_dir), None, str(output_path))
        except Exception as e:
            st.error(f"Something went wrong: {e}")
        else:
            xl = pd.ExcelFile(output_path)
            summary_df = xl.parse("Summary")
            passed_df = xl.parse("Passed - Client Ready")
            failed_df = xl.parse("Failed - Internal Detail")
            supplier_df = xl.parse("Supplier Rejection Report")
            manual_email_df = xl.parse("Needs Manual Email Check")

            # --- AI-assisted TAL alias matching: reclassify leads that failed only
            #     because of a domain/name variant of a real TAL company ---
            if use_ai_tal_alias_matching and tal_file:
                if not anthropic_key:
                    st.warning("Add your Anthropic API key in the sidebar to use AI TAL alias matching. Skipped for this run.")
                else:
                    tal_only_fails = failed_df[
                        failed_df["reasons"].str.contains("not found on TAL", na=False)
                        & ~failed_df["reasons"].str.contains(";")  # only reclassify leads whose SOLE failure was TAL
                    ]
                    if not tal_only_fails.empty:
                        try:
                            reloaded_tal = load_tal(str(tmp / tal_file.name), config["tal"].get("sheet"))
                            name_col = next((c for c in reloaded_tal.columns
                                             if c.lower().strip() in ("company name", "company", "account name", "account")), None)
                            tal_entries = [
                                {"company": r[name_col] if name_col else "", "domain": r["_domain_norm"]}
                                for _, r in reloaded_tal.iterrows()
                            ]
                            failed_leads_payload = [
                                {"failed_domain": r["domain"], "failed_company": r["company"]}
                                for _, r in tal_only_fails.drop_duplicates(subset=["domain"]).iterrows()
                            ]
                            with st.spinner(f"Checking {len(failed_leads_payload)} TAL-rejected companies against real-world knowledge..."):
                                matches = classify_tal_aliases_with_ai(failed_leads_payload, tal_entries, anthropic_key, anthropic_workspace_id)
                            if matches:
                                matched_domains = {m["failed_domain"]: m for m in matches}
                                st.info(
                                    "AI matched these to real TAL companies under a different domain/name: " +
                                    "; ".join(f"{d} → {m['matched_tal_company']} ({m['matched_tal_domain']})" for d, m in matched_domains.items())
                                )
                                move_mask = failed_df["domain"].isin(matched_domains.keys()) & failed_df["reasons"].str.contains("not found on TAL", na=False) & ~failed_df["reasons"].str.contains(";")
                                rows_to_move = failed_df[move_mask].copy()
                                rows_to_move["tal_alias_note"] = rows_to_move["domain"].map(lambda d: f"AI-matched TAL alias: {matched_domains[d]['matched_tal_company']} ({matched_domains[d]['matched_tal_domain']})")
                                rows_to_move_clean = rows_to_move.drop(columns=["status", "reasons", "rejection_category", "needs_manual_qa"])
                                passed_df = pd.concat([passed_df, rows_to_move_clean], ignore_index=True)
                                failed_df = failed_df[~move_mask]
                                supplier_df = supplier_df[~supplier_df["domain"].isin(matched_domains.keys())]
                                summary_df.loc[summary_df["Metric"] == "Passed this run", "Value"] = len(passed_df)
                                summary_df.loc[summary_df["Metric"] == "Failed this run", "Value"] = len(failed_df)
                                with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                                    summary_df.to_excel(writer, sheet_name="Summary", index=False)
                                    passed_df.to_excel(writer, sheet_name="Passed - Client Ready", index=False)
                                    failed_df.to_excel(writer, sheet_name="Failed - Internal Detail", index=False)
                                    supplier_df.to_excel(writer, sheet_name="Supplier Rejection Report", index=False)
                            else:
                                st.caption("AI TAL alias matching found no confident matches among the TAL-rejected leads.")
                        except Exception as e:
                            st.error(f"AI TAL alias matching failed, continuing with standard results only: {e}")

            st.success(f"Done — {len(passed_df)} passed, {len(failed_df)} failed.")
            st.dataframe(summary_df, use_container_width=True, hide_index=True)

            st.download_button("Download full result workbook", output_path.read_bytes(),
                                file_name=f"validated_{leads_file.name.rsplit('.', 1)[0]}.xlsx")

            if not manual_email_df.empty:
                st.download_button(
                    f"Download catch-all emails to test ({len(manual_email_df)} leads)",
                    manual_email_df.to_csv(index=False),
                    file_name=f"catchall_to_test_{leads_file.name.rsplit('.', 1)[0]}.csv",
                    mime="text/csv",
                )

            with st.expander("Preview: Passed leads (client-ready)"):
                st.dataframe(passed_df, use_container_width=True)
            with st.expander("Preview: Failed leads + reasons"):
                st.dataframe(failed_df, use_container_width=True)
            with st.expander("Preview: Supplier rejection report"):
                st.dataframe(supplier_df, use_container_width=True)
            if not manual_email_df.empty:
                with st.expander(f"Preview: Catch-all emails needing manual test ({len(manual_email_df)})"):
                    st.dataframe(manual_email_df, use_container_width=True)

            if mode == "Build a new one":
                st.download_button("Download this config to reuse next time", yaml.dump(config),
                                    file_name=f"{config['campaign']['name'].replace(' ', '_')}_config.yaml")
